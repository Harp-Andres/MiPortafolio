"""Excel workbook generator for CV"""

from datetime import datetime
from pathlib import Path
from typing import Dict, Any
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.domain.entities.cv_models import CVDataModel


class ExcelGenerator:
    """Generate structured Excel workbook from CV data"""

    def __init__(self, cv_data: CVDataModel):
        """Initialize generator with CV data"""
        self.cv_data = cv_data
        self.wb = Workbook()
        self.ws = self.wb.active
        self.ws.title = "Profile"
        
        # Define styles
        self.header_fill = PatternFill(start_color="0066CC", end_color="0066CC", fill_type="solid")
        self.header_font = Font(color="FFFFFF", bold=True, size=12)
        self.subheader_fill = PatternFill(start_color="E8F0FF", end_color="E8F0FF", fill_type="solid")
        self.subheader_font = Font(bold=True, size=11)
        self.border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

    def _style_cell(self, cell, font=None, fill=None, alignment=None, border=True):
        """Apply styles to a cell"""
        if font:
            cell.font = font
        if fill:
            cell.fill = fill
        if alignment:
            cell.alignment = alignment
        if border:
            cell.border = self.border

    def create_profile_sheet(self) -> None:
        """Create profile information sheet"""
        ws = self.ws
        profile = self.cv_data.profile

        # Title
        ws['A1'] = "PROFESSIONAL PROFILE"
        self._style_cell(ws['A1'], font=self.header_font, fill=self.header_fill)
        ws.merge_cells('A1:D1')

        # Profile data
        row = 3
        data = [
            ("Name", profile.name),
            ("Title", profile.title),
            ("Email", profile.email),
            ("Phone", profile.phone or "N/A"),
            ("Location", profile.location),
            ("GitHub", profile.github),
            ("LinkedIn", profile.linkedin),
            ("Portfolio", profile.portfolio),
        ]

        for label, value in data:
            ws[f'A{row}'] = label
            ws[f'B{row}'] = value
            self._style_cell(ws[f'A{row}'], font=Font(bold=True))
            self._style_cell(ws[f'B{row}'])
            row += 1

        # Professional summary
        ws[f'A{row + 1}'] = "PROFESSIONAL SUMMARY"
        self._style_cell(ws[f'A{row + 1}'], font=self.subheader_font, fill=self.subheader_fill)
        ws.merge_cells(f'A{row + 1}:D{row + 1}')

        ws[f'A{row + 2}'] = profile.bio
        ws[f'A{row + 2}'].alignment = Alignment(wrap_text=True, vertical='top')
        self._style_cell(ws[f'A{row + 2}'])

        # Adjust column widths
        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 50
        ws.column_dimensions['C'].width = 30

    def create_skills_sheet(self) -> None:
        """Create skills sheet"""
        if not self.cv_data.skills:
            return

        ws = self.wb.create_sheet("Skills")

        # Header
        ws['A1'] = "TECHNICAL SKILLS"
        self._style_cell(ws['A1'], font=self.header_font, fill=self.header_fill)
        ws.merge_cells('A1:B1')

        row = 3
        ws['A3'] = "Category"
        ws['B3'] = "Skills"
        for col in ['A', 'B']:
            self._style_cell(ws[f'{col}3'], font=self.subheader_font, fill=self.subheader_fill)

        row = 4
        for category in self.cv_data.skills:
            ws[f'A{row}'] = category.category
            skills_str = ", ".join([s.name for s in category.skills])
            ws[f'B{row}'] = skills_str
            ws[f'B{row}'].alignment = Alignment(wrap_text=True)
            self._style_cell(ws[f'A{row}'])
            self._style_cell(ws[f'B{row}'])
            row += 1

        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 70

    def create_experience_sheet(self) -> None:
        """Create experience sheet"""
        if not self.cv_data.experience:
            return

        ws = self.wb.create_sheet("Experience")

        # Header
        ws['A1'] = "PROFESSIONAL EXPERIENCE"
        self._style_cell(ws['A1'], font=self.header_font, fill=self.header_fill)
        ws.merge_cells('A1:E1')

        row = 3
        headers = ['Company', 'Position', 'Period', 'Technologies', 'Description']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col)
            cell.value = header
            self._style_cell(cell, font=self.subheader_font, fill=self.subheader_fill)

        row = 4
        for exp in self.cv_data.experience:
            ws[f'A{row}'] = exp.company
            ws[f'B{row}'] = exp.title
            ws[f'C{row}'] = exp.period
            ws[f'D{row}'] = ", ".join(exp.technologies) if exp.technologies else ""
            ws[f'E{row}'] = exp.description

            for col in range(1, 6):
                cell = ws.cell(row=row, column=col)
                cell.alignment = Alignment(wrap_text=True, vertical='top')
                self._style_cell(cell)

            row += 1

        # Set column widths
        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 25
        ws.column_dimensions['C'].width = 20
        ws.column_dimensions['D'].width = 30
        ws.column_dimensions['E'].width = 50

        # Set row heights for wrapped text
        for r in range(4, row):
            ws.row_dimensions[r].height = 50

    def create_education_sheet(self) -> None:
        """Create education sheet"""
        if not self.cv_data.education:
            return

        ws = self.wb.create_sheet("Education")

        # Header
        ws['A1'] = "EDUCATION"
        self._style_cell(ws['A1'], font=self.header_font, fill=self.header_fill)
        ws.merge_cells('A1:D1')

        row = 3
        headers = ['Degree', 'Institution', 'Graduation Year', 'Description']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col)
            cell.value = header
            self._style_cell(cell, font=self.subheader_font, fill=self.subheader_fill)

        row = 4
        for edu in self.cv_data.education:
            ws[f'A{row}'] = edu.degree
            ws[f'B{row}'] = edu.institution
            ws[f'C{row}'] = edu.graduation
            ws[f'D{row}'] = edu.description or ""

            for col in range(1, 5):
                cell = ws.cell(row=row, column=col)
                cell.alignment = Alignment(wrap_text=True)
                self._style_cell(cell)

            row += 1

        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 35
        ws.column_dimensions['C'].width = 15
        ws.column_dimensions['D'].width = 40

    def create_certificates_sheet(self) -> None:
        """Create certificates sheet"""
        if not self.cv_data.certificates:
            return

        ws = self.wb.create_sheet("Certificates")

        # Header
        ws['A1'] = "CERTIFICATIONS & TRAINING"
        self._style_cell(ws['A1'], font=self.header_font, fill=self.header_fill)
        ws.merge_cells('A1:D1')

        row = 3
        headers = ['Category', 'Certificate Name', 'Issuer', 'Hours']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col)
            cell.value = header
            self._style_cell(cell, font=self.subheader_font, fill=self.subheader_fill)

        row = 4
        for category, certs in self.cv_data.certificates.items():
            for cert in certs:
                ws[f'A{row}'] = category
                ws[f'B{row}'] = cert.name
                ws[f'C{row}'] = cert.issuer
                ws[f'D{row}'] = cert.hours or ""

                for col in range(1, 5):
                    cell = ws.cell(row=row, column=col)
                    self._style_cell(cell)

                row += 1

        ws.column_dimensions['A'].width = 25
        ws.column_dimensions['B'].width = 50
        ws.column_dimensions['C'].width = 25
        ws.column_dimensions['D'].width = 10

    def generate(self, output_path: Path) -> Path:
        """Generate complete Excel workbook"""
        self.create_profile_sheet()
        self.create_skills_sheet()
        self.create_experience_sheet()
        self.create_education_sheet()
        self.create_certificates_sheet()

        self.wb.save(str(output_path))
        return output_path

    @staticmethod
    def from_dict(data: Dict[str, Any]) -> 'ExcelGenerator':
        """Create generator from dictionary (from cv-data.ts)"""
        cv_model = CVDataModel(**data)
        return ExcelGenerator(cv_model)

